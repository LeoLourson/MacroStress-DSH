"""Bank of Russia banking statistics: credit debt, the mortgage rate and overdue debt by activity.

Recent months come from the open-data service (``/dataservice/data``); earlier ones from the
Excel export of the archive tables of www.cbr.ru (``/Queries/StatTable/Excel/<table>``). The
overdue share by activity comes from a workbook of the statistics pages (``/vfs/statistics``).
Where both have a month, the service wins. Both date a value by the first day after its
month: a stock "на 01.02" and the rate of loans issued in January belong to January. No key
is needed.
"""

import io
from datetime import date, datetime
from typing import Literal

import httpx
import openpyxl

from mst.contracts import Contract

from .release import Release

SERVICE_URL = "https://www.cbr.ru/dataservice/data"
EXCEL_URL = "https://www.cbr.ru/Queries/StatTable/Excel/"
DEBT_BY_ACTIVITY = (
    "https://www.cbr.ru/vfs/statistics/BankSector/Loans_to_corporations/"
    "01_02_A_Debt_corp_by_activity.xlsx"
)


def _previous_month(day: date) -> date:
    return date(day.year - 1, 12, 1) if day.month == 1 else date(day.year, day.month - 1, 1)


def _next_month(month: date) -> date:
    return date(month.year + 1, 1, 1) if month.month == 12 else date(month.year, month.month + 1, 1)


def _service(
    http: httpx.Client, start: date, publication: int, dataset: int, element: int, measure=None
) -> dict[date, float]:
    params = {
        "y1": start.year,
        "y2": date.today().year,
        "publicationId": publication,
        "datasetId": dataset,
    }
    if measure is not None:
        params["measureId"] = measure
    response = http.get(SERVICE_URL, params=params)
    response.raise_for_status()
    return {
        _previous_month(date.fromisoformat(row["date"][:10])): float(row["obs_val"])
        for row in response.json()["RawData"]
        if row["element_id"] == element
    }


def _excel(http: httpx.Client, table: str, day: date | None = None) -> list[tuple]:
    """The rows of the first sheet of an archive table, read in memory."""
    params = {"lang": "ru-RU"}
    if day is not None:
        params["dt"] = day.strftime("%Y%m%d")
    response = http.get(EXCEL_URL + table, params=params)
    response.raise_for_status()
    book = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    return list(book.worksheets[0].iter_rows(values_only=True))


class Group(Contract):
    """One group of borrowers: its dataset in the service and its archive table."""

    publication: int
    dataset: int
    measure: int
    """The region of the service, 22 for the whole of Russia."""
    archive: str
    """The archive table with the debt on the first day of each month, e.g. 302-01M."""


class CbrCredit(Release):
    """Debt on credits at the end of the month, mln roubles, summed over ``groups``."""

    loader: Literal["cbr_credit"]
    element: int
    """The service column: 35 for the debt in all currencies."""
    groups: list[Group]

    def fetch(self, http: httpx.Client, start: date) -> dict[date, float]:
        parts = []
        for group in self.groups:
            values = _service(
                http, start, group.publication, group.dataset, self.element, group.measure
            )
            month, first = start, min(values, default=start)
            while month < first:
                values[month] = self._archive(http, group.archive, _next_month(month))
                month = _next_month(month)
            parts.append(values)
        common = set.intersection(*(set(part) for part in parts))
        return {month: sum(part[month] for part in parts) for month in sorted(common)}

    @staticmethod
    def _archive(http: httpx.Client, table: str, day: date) -> float:
        """The total debt for Russia on ``day``, mln roubles, from the archive table."""
        rows = [list(row) for row in _excel(http, table, day)]
        cells = [str(cell) for row in rows for cell in row if cell is not None]
        header = next((row for row in rows if row[1:2] == ["Объем"]), None)
        total = next((row for row in rows if row[:1] == ["РОССИЙСКАЯ ФЕДЕРАЦИЯ"]), None)
        if (
            "млн.руб." not in cells
            or not any(cell.endswith(f"на {day:%d.%m.%Y}") for cell in cells)
            or header is None
            or header[4:5] != ["Задолженность"]
            or total is None
            or not isinstance(total[4], int | float)
        ):
            raise ValueError(f"table {table} on {day} has no debt of the Russian Federation")
        # Volume total, roubles, currency; then the debt total.
        return float(total[4])


class CbrMortgageRate(Release):
    """Weighted average rate of mortgage loans issued in the month, % per year."""

    loader: Literal["cbr_mortgage_rate"]
    publication: int
    dataset: int
    element: int
    """The service column: 36 for loans in roubles."""
    archive: str
    """The archive table with the monthly rows, 4-3 for loans in roubles."""

    def fetch(self, http: httpx.Client, start: date) -> dict[date, float]:
        values = self._archive(http)
        values |= _service(http, start, self.publication, self.dataset, self.element)
        return {month: value for month, value in sorted(values.items()) if month >= start}

    def _archive(self, http: httpx.Client) -> dict[date, float]:
        rows = _excel(http, self.archive)
        if (
            len(rows) < 4
            or len(rows[2]) < 8
            or not str(rows[2][7]).startswith("Средневзвешенная ставка")
            or len(rows[3]) < 9
            or rows[3][8] != "по кредитам, выданным в течение месяца"
        ):
            raise ValueError(f"table {self.archive} has no rate of loans issued in the month")
        # Date, count, volume, debt and overdue, two terms, then the two rates.
        return {
            _previous_month(row[0].date()): float(row[8])
            for row in rows[4:]
            if isinstance(row[0], datetime) and isinstance(row[8], int | float)
        }


def _activity(label) -> str:
    """A row label without its indent and the trailing «, из них:» of a row with parts."""
    return str(label).strip().removesuffix(", из них:")


class CbrOverdue(Release):
    """The overdue share of the debt on credits to organisations of one activity, a fraction.

    Table 01_02_A gives the debt (sheet «итого») and the overdue debt (sheet «в т.ч.
    просроч.») by activity on the first day of each month since 01.02.2019, mln roubles, all
    currencies.
    """

    loader: Literal["cbr_overdue"]
    activity: str
    """The row label, e.g. «обрабатывающие производства» or «ВСЕГО»."""

    def fetch(self, http: httpx.Client, start: date) -> dict[date, float]:
        response = http.get(DEBT_BY_ACTIVITY)
        response.raise_for_status()
        book = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        sheets = {sheet.title.strip(): list(sheet.iter_rows(values_only=True)) for sheet in book}
        debt, overdue = (self._row(sheets, name) for name in ("итого", "в т.ч. просроч."))
        return {
            month: overdue[month] / debt[month]
            for month in sorted(debt)
            if month >= start and month in overdue and debt[month] > 0
        }

    def _row(self, sheets: dict, name: str) -> dict[date, float]:
        rows = sheets.get(name)
        if not rows or len(rows) < 3 or not str(rows[0][0]).endswith("млн руб."):
            raise ValueError(f"table 01_02_A has no sheet «{name}» in mln roubles")
        found = [row for row in rows[2:] if _activity(row[0]) == self.activity]
        if len(found) != 1:
            raise ValueError(f"sheet «{name}» has {len(found)} rows «{self.activity}»")
        return {
            _previous_month(day.date()): float(value)
            for day, value in zip(rows[1][1:], found[0][1:], strict=False)
            if isinstance(day, datetime) and isinstance(value, int | float)
        }
