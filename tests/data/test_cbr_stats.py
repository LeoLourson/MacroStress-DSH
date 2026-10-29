import io
import json
from datetime import date, datetime

import httpx
import openpyxl
import pytest

from mst.data.cbr_stats import CbrCredit, CbrMortgageRate, CbrOverdue, Group


def service(rows):
    """A dataservice answer: one record per (date, element)."""
    raw = [
        {"element_id": element, "measure_id": 22, "obs_val": value, "date": f"{day}T00:00:00"}
        for day, element, value in rows
    ]
    return httpx.Response(200, text=json.dumps({"RawData": raw}))


def workbook(rows):
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return httpx.Response(200, content=buffer.getvalue())


def debt_table(day, total):
    """The archive table 302-01M or 302-02M on ``day``, as the Bank exports it."""
    return workbook(
        [
            ["Информация о кредитах, предоставленных юридическим лицам *"],
            ["млн.руб."],
            [],
            ["", f"Кредиты, предоставленные юридическим лицам по состоянию  на {day}"],
            [None, "Объем", None, None, "Задолженность"],
            [None, "Всего", "в рублях", "в иностранной валюте", "Всего", "в рублях"],
            ["РОССИЙСКАЯ ФЕДЕРАЦИЯ", 1, 1, 0, total, 1, 0, 5],
            ["ЦЕНТРАЛЬНЫЙ ФЕДЕРАЛЬНЫЙ ОКРУГ", 1, 1, 0, 7, 1, 0, 5],
        ]
    )


def client(respond, seen):
    def record(request):
        seen.append(request)
        return respond(request)

    return httpx.Client(transport=httpx.MockTransport(record))


CREDIT = CbrCredit(
    loader="cbr_credit",
    lag_days=38,
    revised=True,
    element=35,
    groups=[
        Group(publication=22, dataset=50, measure=22, archive="302-01M"),
        Group(publication=20, dataset=42, measure=22, archive="302-02M"),
    ],
)


def archive_day(request):
    return datetime.strptime(request.url.params["dt"], "%Y%m%d").strftime("%d.%m.%Y")


def test_credit_sums_the_groups_and_takes_the_archive_before_the_service():
    def respond(request):
        if request.url.path == "/dataservice/data":
            base = 100.0 if request.url.params["datasetId"] == "50" else 10.0
            return service([("2012-04-01", 35, base + 3), ("2012-04-01", 36, 1.0)])
        corporate = request.url.path.endswith("/302-01M")
        totals = {"01.02.2012": (1000, 100), "01.03.2012": (2000, 200)}
        day = archive_day(request)
        return debt_table(day, totals[day][0 if corporate else 1])

    seen = []
    values = CREDIT.fetch(client(respond, seen), date(2012, 1, 1))
    assert values == {
        date(2012, 1, 1): 1100.0,
        date(2012, 2, 1): 2200.0,
        date(2012, 3, 1): 116.0,
    }
    archive = sorted(
        (r.url.path, r.url.params["dt"]) for r in seen if r.url.path.startswith("/Queries/")
    )
    assert archive == [
        ("/Queries/StatTable/Excel/302-01M", "20120201"),
        ("/Queries/StatTable/Excel/302-01M", "20120301"),
        ("/Queries/StatTable/Excel/302-02M", "20120201"),
        ("/Queries/StatTable/Excel/302-02M", "20120301"),
    ]
    data = [r for r in seen if r.url.path == "/dataservice/data"][0]
    assert (data.url.params["measureId"], data.url.params["y1"]) == ("22", "2012")


def test_credit_refuses_an_archive_table_of_another_day():
    def respond(request):
        if request.url.path == "/dataservice/data":
            return service([("2012-03-01", 35, 1.0)])
        return debt_table("01.01.2019", 5)

    with pytest.raises(ValueError, match="302-01M"):
        CREDIT.fetch(client(respond, []), date(2012, 1, 1))


def test_credit_refuses_an_archive_table_it_cannot_read():
    def respond(request):
        if request.url.path == "/dataservice/data":
            return service([("2012-03-01", 35, 1.0)])
        return workbook([["нет данных"]])

    with pytest.raises(ValueError, match="302-01M"):
        CREDIT.fetch(client(respond, []), date(2012, 1, 1))


MORTGAGE = CbrMortgageRate(
    loader="cbr_mortgage_rate",
    lag_days=31,
    revised=True,
    publication=14,
    dataset=29,
    element=36,
    archive="4-3",
)


def mortgage_table():
    """The archive table 4-3: one row per first day, the rate of the month in column 8."""
    return workbook(
        [
            ["Ипотечные жилищные кредиты в рублях"],
            [],
            [
                "",
                "Количество",
                "Объем",
                "Задолженность",
                None,
                "Срок",
                None,
                "Средневзвешенная ставка, %",
                None,
                "Справочно",
            ],
            [
                None,
                None,
                None,
                "Всего",
                "В том числе просроченная",
                "по выданным с начала года кредитам",
                "по кредитам, выданным в течение месяца",
                "по выданным с начала года кредитам",
                "по кредитам, выданным в течение месяца",
            ],
            [datetime(2011, 10, 1), "X", 1, 2, 3, 4.0, "X", 13.9, "X", "X"],
            [datetime(2012, 2, 1), 1, 2, 3, 4, 5.0, 6.0, 12.1, 12.2, 9],
            [datetime(2012, 3, 1), 1, 2, 3, 4, 5.0, 6.0, 12.0, 11.9, 9],
        ]
    )


def test_mortgage_rate_reads_the_month_rate_and_prefers_the_service():
    seen = []

    def respond(request):
        if request.url.path == "/dataservice/data":
            return service([("2012-03-01", 36, 11.5), ("2012-03-01", 40, 0.0)])
        return mortgage_table()

    values = MORTGAGE.fetch(client(respond, seen), date(2012, 1, 1))
    assert values == {date(2012, 1, 1): 12.2, date(2012, 2, 1): 11.5}
    assert any(r.url.path == "/Queries/StatTable/Excel/4-3" for r in seen)


def test_mortgage_rate_refuses_a_table_without_the_month_rate():
    def respond(request):
        if request.url.path == "/dataservice/data":
            return service([])
        return workbook([["Ипотечные жилищные кредиты"], [], ["", "Количество"]])

    with pytest.raises(ValueError, match="4-3"):
        MORTGAGE.fetch(client(respond, []), date(2012, 1, 1))


def by_activity(sheets):
    """Table 01_02_A: ``sheets[name]`` is a list of (label, value per day) rows."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    days = [datetime(2019, 2, 1), datetime(2019, 3, 1)]
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(["Задолженность по кредитам ..., млн руб."])
        sheet.append(["Виды экономической деятельности", *days])
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def overdue_client(content):
    return httpx.Client(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, content=content))
    )


def test_overdue_share_divides_the_overdue_row_by_the_debt_row_of_the_activity():
    content = by_activity(
        {
            "итого": [("ВСЕГО", 1000, 1200), ("строительство, из них:", 100, 120)],
            "в т.ч. просроч.": [("ВСЕГО", 50, 60), ("строительство, из них:", 20, 30)],
        }
    )
    source = CbrOverdue(loader="cbr_overdue", activity="строительство", lag_days=37, revised=True)
    values = source.fetch(overdue_client(content), date(2012, 1, 1))
    assert values == {date(2019, 1, 1): pytest.approx(0.2), date(2019, 2, 1): pytest.approx(0.25)}


def test_overdue_share_refuses_a_table_without_the_activity():
    content = by_activity({"итого": [("ВСЕГО", 1, 1)], "в т.ч. просроч.": [("ВСЕГО", 1, 1)]})
    source = CbrOverdue(loader="cbr_overdue", activity="строительство", lag_days=37, revised=True)
    with pytest.raises(ValueError, match="0 rows"):
        source.fetch(overdue_client(content), date(2012, 1, 1))
